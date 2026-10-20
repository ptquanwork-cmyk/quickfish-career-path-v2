"""Build the career records in public/data/career-dashboard.v1.json from ALLJOBs_slim.xlsx only.

The whole file is rebuilt from the sheet; nothing from earlier data is kept.
Fields the sheet does not have are left empty (null / []) so the dashboard shows
"no data" or hides the tile instead of older values.

Usage: python scripts/import-alljobs-sheet.py path/to/ALLJOBs_slim.xlsx
"""
import collections
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
JSON_PATH = ROOT / "public" / "data" / "career-dashboard.v1.json"
SHEET = Path(sys.argv[1])
SOURCE = "ALLJOBs"
REVIEWED = "md-reviewed"
STAGE_LEVELS = ["Entry", "Mid", "Senior", "Executive"]
# ponytail: sheet has 4 stage-skill groups for 5 path rungs; Executive goes to the top rung, rung 4 keeps no stage skills.
STAGE_TO_RUNG = [0, 1, 2, 4]
TRENDS = {"strong-growth", "top-growth", "growth", "stable", "decline"}


def split(value, sep=","):
    return [part.strip().removeprefix("and ") for part in str(value or "").split(sep) if part.strip()]


def stage(title, level=None, skills=(), provenance=REVIEWED):
    return {"level": level, "title": title, "description": None, "skills": list(skills), "dayToDay": None,
            "salaryUs": None, "salaryVietnam": None, "outlook": None, "provenance": provenance}


rows = list(load_workbook(SHEET, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
headers = rows[0]
sheet = [dict(zip(headers, row)) for row in rows[1:] if row[0] is not None]
slugs = {r["slug"] for r in sheet}
if len(slugs) != len(sheet):
    sys.exit("Duplicate slugs in sheet")

careers = []
for r in sheet:
    if r["trend_status"] not in TRENDS:
        sys.exit(f"{r['slug']}: unknown trend_status {r['trend_status']!r}")
    path = [r[f"path_{i}"] for i in range(1, 6) if r[f"path_{i}"]]
    stages = [split(group.rstrip(".")) for group in str(r["stage_skills"]).split("., ")]
    legacy = [stage(title, provenance="unavailable") for title in path]
    for level, skills, rung in zip(STAGE_LEVELS, stages, STAGE_TO_RUNG):
        if rung < len(path):
            legacy[rung] = stage(path[rung], level, skills)
    departments = split(r["departments"])
    majors = [r[f"major_{i}"] for i in range(1, 4) if r[f"major_{i}"]]
    similar = split(r["similar_careers"], "|")
    if unknown := [s for s in similar if s not in slugs]:
        sys.exit(f"{r['slug']}: unknown similar careers {unknown}")

    careers.append({
        "id": int(r["id"]),
        "slug": r["slug"],
        "name": {"en": r["name_en"], "vi": r["name_vi"]},
        "jobGroup": r["category_job_group"],
        "minimumExperienceYears": None,
        "riasec": {"code": r["riasec_code"], "scores": {k: float(r[f"riasec_{k}"]) for k in "RIASEC"}},
        "salary": {"vietnam": r["salary_vietnam"], "unitedStates": r["salary_united_states"], "provenance": REVIEWED},
        "description": None,
        "skillSignals": [{"name": r[f"skill_{i}"], "percent": int(r[f"skill_{i}_employers_pct"]), "scope": "industry",
                          "source": SOURCE, "growth": None} for i in range(1, 9) if r[f"skill_{i}"]],
        "industrySignal": None,
        "majors": majors,
        "education": [{
            "name": major,
            # The sheet lists departments per career, not per major; only a single department is unambiguous.
            "department": departments[0] if len(departments) == 1 else None,
            "coreSubjects": split(r["core_subjects"]),
            "examBlocks": split(r["exam_blocks"]),
            "provenance": REVIEWED,
        } for major in majors],
        "departments": departments,
        "careerPath": path,
        "skills": list(dict.fromkeys(skill for group in stages for skill in group)),
        "dayToDay": r["day_to_day"],
        "legacyStages": legacy,
        "trend": {"rank": int(r["popularity_rank"]), "globalNetGrowth": None, "wefMentions": None, "vietnamNetGrowth": None,
                  "status": r["trend_status"], "sourceLabel": SOURCE, "provenance": REVIEWED},
        "aiImpact": {"level": "unassessed", "summary": "The ALLJOBs sheet has no occupation-level AI impact data.",
                     "horizon": "2025–2030", "provenance": "unavailable"},
        "sourceTier": REVIEWED,
        "evidence": [{"section": SOURCE, "label": SHEET.name, "provenance": REVIEWED}],
        "similarCareerSlugs": similar,
    })

careers.sort(key=lambda c: c["id"])
major_careers = collections.defaultdict(list)
for c in careers:
    for major in c["majors"]:
        major_careers[major].append(c["name"]["en"])
job_groups = sorted({c["jobGroup"] for c in careers})
data = {
    "metadata": {
        "version": "2.0.0",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFile": SHEET.name,
        "totalCareers": len(careers),
        "totalJobGroups": len(job_groups),
        "totalMajors": len(major_careers),
        "trendCounts": dict(collections.Counter(c["trend"]["status"] for c in careers)),
        "notes": [f"Everything comes from {SHEET.name}; fields the sheet lacks are empty."],
    },
    "jobGroups": job_groups,
    "majors": [{"name": name, "careerNames": names} for name, names in sorted(major_careers.items())],
    "careers": careers,
}
JSON_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Wrote {len(careers)} careers from {SHEET.name}; trend counts {data['metadata']['trendCounts']}")
