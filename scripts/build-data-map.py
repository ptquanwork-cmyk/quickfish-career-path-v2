"""Build data/QuickFish_CareerPath_V2_DataMap.xlsx — which data feeds which part of the V2 UI.

Sources (nothing is invented):
  * data/ALLJOBs_slim.xlsx                       — the career sheet (132 rows)
  * src/features/career-dashboard/report-recommendations.ts — report-name crosswalk
  * the V2 code itself, described row by row in UI_MAP below

Usage: python scripts/build-data-map.py
"""
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SHEET = ROOT / "data" / "ALLJOBs_slim.xlsx"
CROSSWALK = ROOT / "src" / "features" / "career-dashboard" / "report-recommendations.ts"
OUT = ROOT / "data" / "QuickFish_CareerPath_V2_DataMap.xlsx"
EXAMPLE_SLUG = "social-scientists-and-related-workers"

FONT = "Arial"
INK = "192332"
PLUM = "A44E80"
GRID = Side(style="thin", color="E4DCE8")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
WRAP = Alignment(wrap_text=True, vertical="top")
SECTION_FILL = {
    "Identity & URL": "EDE7F6", "Hero (header)": "D6EEF8", "Metric strip": "FDE7F1", "Filters & search": "EAF4EA",
    "Overview tab": "FFF3D6", "Skills tab": "FFF3D6", "Job opportunities tab": "E8F0FB", "Majors tab": "F3E8FF",
    "Path tab": "E6F6F3", "Related careers": "F1F1F1", "Right rail": "FDECEC", "Mobile card": "FFF8E1",
    "Backend (per user)": "ECEFF4",
}

rows = list(load_workbook(SHEET, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
headers = list(rows[0])
careers = [dict(zip(headers, r)) for r in rows[1:] if r[0] is not None]
by_slug = {c["slug"]: c for c in careers}
ex = by_slug[EXAMPLE_SLUG]


def stage_groups(value):
    return [g.strip().rstrip(".") for g in str(value or "").split("., ")]


def vnd(value):
    match = re.match(r"^(.*?)\s*million VND/month$", str(value or ""), re.I)
    return f"{match.group(1)} triệu · VNĐ / tháng" if match else str(value or "")


TREND_VI = {"strong-growth": "Tăng mạnh", "top-growth": "Top tăng", "growth": "Đang tăng", "stable": "Ổn định", "decline": "Đang giảm"}
RIASEC_VI = {"R": "Thực tế", "I": "Nghiên cứu", "A": "Sáng tạo", "S": "Xã hội", "E": "Lãnh đạo", "C": "Quy trình"}
DASH = "src/features/career-dashboard/"

# (screen, area, UI element, example, source, sheet columns, rule, JSON field, code)
UI_MAP = [
    ("Web + mobile", "Identity & URL", "Career id shown as #127", f"#{ex['id']:03d}", "Sheet", "id", "Zero-padded to 3 digits", "careers[].id", DASH + "CareerDashboard.tsx → conceptHero2"),
    ("Web + mobile", "Identity & URL", "?career=… in the URL, saved-career key", ex["slug"], "Sheet", "slug", "Unique key; used for links, bookmarks and similar-career references", "careers[].slug", DASH + "CareerDashboard.tsx → selectedCareer / updateQuery"),
    ("Web", "Hero (header)", "Eyebrow '#127 · Social Sciences & Education'", f"#{ex['id']:03d} · {ex['category_job_group']}", "Sheet", "id, category_job_group", "", "careers[].jobGroup", DASH + "CareerDashboard.tsx → conceptHero2"),
    ("Web", "Hero (header)", "Big title", ex["name_vi"], "Sheet", "name_vi (VI) / name_en (EN)", "Language follows the app locale", "careers[].name.vi / .en", DASH + "CareerDashboard.tsx → careerName()"),
    ("Web", "Hero (header)", "One-line description under the title", ex["day_to_day"], "Sheet", "day_to_day", "The sheet has it in English only; shown in both languages", "careers[].dayToDay", DASH + "career-summary.ts → getCareerSummary()"),
    ("Web", "Hero (header)", "RIASEC pill", f"RIASEC {ex['riasec_code']}", "Sheet", "riasec_code", "", "careers[].riasec.code", DASH + "CareerDashboard.tsx → conceptHero2"),
    ("Web", "Hero (header)", "Sticker illustration (bottom right)", f"/assets/career-deck/{{group}}.png", "Sheet", "category_job_group", "One PNG per job group (9 files)", "careers[].jobGroup", DASH + "MobileCareerDeck.tsx → careerCover()"),
    ("Web", "Hero (header)", "Mascot + caption 'Smarty Salmon · INT'", "salmon.png", "Backend", "— (user survey)", "latest_survey.fish_type → fish theme (colour + image); fish menu can preview others", "—", DASH + "design/fish-theme.ts → getFishTheme()"),
    ("Web", "Hero (header)", "Ocean background", "/assets/career-path-v3/ocean.webp", "Asset", "—", "Static image + white veil for legibility", "—", DASH + "career-versions.css → .qf-c2-hero"),
    ("Web", "Hero (header)", "'Báo cáo chi tiết' button", "opens PDF", "Backend", "— (user survey)", "POST /api/survey/report/generate/<latest_survey.id>/ {language} → pdf_url, opened in a new tab", "—", DASH + "CareerDashboard.tsx → openDetailReport()"),
    ("Web", "Metric strip", "Lương khởi điểm", vnd(ex["salary_vietnam"]), "Sheet", "salary_vietnam", "'13–30 million VND/month' split into '13–30 triệu' + 'VNĐ / tháng'", "careers[].salary.vietnam", DASH + "CareerDashboard.tsx → salaryParts"),
    ("Web", "Metric strip", "Hạng phổ biến '#52 / 132'", f"#{ex['popularity_rank']} / {len(careers)}", "Sheet", "popularity_rank", "Denominator = number of rows in the sheet", "careers[].trend.rank", DASH + "CareerDashboard.tsx → c2 strip"),
    ("Web", "Metric strip", "Trend chip under the rank", TREND_VI[ex["trend_status"]], "Sheet", "trend_status", "strong-growth / top-growth / growth / stable / decline → label + colour", "careers[].trend.status", DASH + "CareerDashboard.tsx → TrendBadge"),
    ("Web", "Metric strip", "Kỹ năng nổi bật + '% nhà tuyển dụng'", f"{ex['skill_1']} · {ex['skill_1_employers_pct']}%", "Sheet", "skill_1, skill_1_employers_pct", "", "careers[].skillSignals[0]", DASH + "CareerDashboard.tsx → topSkill"),
    ("Web + mobile", "Filters & search", "Nhóm nghề filter", "9 options", "Sheet", "category_job_group", "Distinct values of the column", "jobGroups[]", DASH + "CareerDashboard.tsx → conceptFilters"),
    ("Web + mobile", "Filters & search", "RIASEC filter", "R I A S E C", "Sheet", "riasec_code", "Career matches if its code contains the letter", "careers[].riasec.code", DASH + "CareerDashboard.tsx → filteredCareers"),
    ("Web + mobile", "Filters & search", "Xu hướng filter", "5 options", "Sheet", "trend_status", "", "careers[].trend.status", DASH + "CareerDashboard.tsx → conceptFilters"),
    ("Web + mobile", "Filters & search", "Search box", "'luật' → Luật sư", "Sheet", "name_vi, name_en, slug", "Case-insensitive substring", "careers[].name / slug", DASH + "CareerDashboard.tsx → filteredCareers"),
    ("Web", "Filters & search", "'Đang khám phá' career selector", ex["name_vi"], "Sheet", "name_vi / name_en", "Lists the careers left after filters", "careers[]", DASH + "CareerDashboard.tsx → careerSelect"),
    ("Web", "Overview tab", "Kỹ năng quan trọng: 3 bars", f"{ex['skill_1']} {ex['skill_1_employers_pct']}% …", "Sheet", "skill_1..3, skill_1..3_employers_pct", "Bar length = % of employers", "careers[].skillSignals[0..2]", DASH + "CareerDashboard.tsx → conceptOverview"),
    ("Web", "Overview tab", "'Xem thêm 5 kỹ năng'", "skill_4..8", "Sheet", "skill_4..8, skill_4..8_employers_pct", "Expands the same list", "careers[].skillSignals[3..7]", DASH + "CareerDashboard.tsx → conceptOverview"),
    ("Web", "Overview tab", "Hồ sơ RIASEC radar — 'Nghề này'", f"R{ex['riasec_R']} I{ex['riasec_I']} A{ex['riasec_A']} S{ex['riasec_S']} E{ex['riasec_E']} C{ex['riasec_C']}", "Sheet", "riasec_R, riasec_I, riasec_A, riasec_S, riasec_E, riasec_C", "Plotted as score / 7 × 100", "careers[].riasec.scores", DASH + "CareerDashboard.tsx → conceptRadar"),
    ("Web", "Overview tab", "Hồ sơ RIASEC radar — 'Hồ sơ của bạn'", "user scores", "Backend", "— (user survey)", "latest_survey.riasec_scores, plotted as score / 5 × 100", "—", DASH + "CareerDashboard.tsx → conceptRadar"),
    ("Web", "Overview tab", "Caption under the radar", " · ".join(RIASEC_VI[k] for k in ex["riasec_code"]), "Sheet", "riasec_code", "Each letter → Vietnamese label", "careers[].riasec.code", DASH + "CareerDashboard.tsx → riasecCaption"),
    ("Web", "Skills tab", "Tất cả kỹ năng (8 bars, 2 columns)", "8 skills", "Sheet", "skill_1..8, skill_1..8_employers_pct", "", "careers[].skillSignals[]", DASH + "CareerDashboard.tsx → conceptSkillsTab"),
    ("Web", "Job opportunities tab", "Hạng phổ biến + trend", f"#{ex['popularity_rank']} · {TREND_VI[ex['trend_status']]}", "Sheet", "popularity_rank, trend_status", "", "careers[].trend", DASH + "CareerDashboard.tsx → conceptMarketTab"),
    ("Web", "Job opportunities tab", "Xu hướng trong nhóm nghề (bars)", "counts per trend", "Sheet (computed)", "category_job_group, trend_status", "Count of careers in the same job group per trend_status; current career's bar highlighted", "careers[]", DASH + "CareerDashboard.tsx → conceptGroupTrends"),
    ("Web", "Job opportunities tab", "Nghề tương tự (chips, clickable)", ex["similar_careers"], "Sheet", "similar_careers", "' | '-separated slugs → name + popularity rank of each", "careers[].similarCareerSlugs", DASH + "CareerDashboard.tsx → conceptSimilar"),
    ("Web", "Majors tab", "One card per major", f"{ex['major_1']}, {ex['major_2']}, {ex['major_3']}", "Sheet", "major_1, major_2, major_3", "", "careers[].majors / education[]", DASH + "CareerDashboard.tsx → conceptEducationTab"),
    ("Web", "Majors tab", "Khối ngành (under the major name)", ex["departments"], "Sheet", "departments", "Given per career, not per major: shown on every card; when the list has 2+ departments all are shown", "careers[].departments", DASH + "CareerDashboard.tsx → conceptEducationTab"),
    ("Web", "Majors tab", "Môn cốt lõi", ex["core_subjects"], "Sheet", "core_subjects", "Comma-separated", "careers[].education[].coreSubjects", DASH + "CareerDashboard.tsx → conceptEducationTab"),
    ("Web", "Majors tab", "Khối xét tuyển chips", ex["exam_blocks"], "Sheet", "exam_blocks", "Comma-separated", "careers[].education[].examBlocks", DASH + "CareerDashboard.tsx → conceptEducationTab"),
    ("Web + mobile", "Path tab", "5 steps of the career path", " → ".join(str(ex[f"path_{i}"]) for i in range(1, 6)), "Sheet", "path_1..path_5", "", "careers[].careerPath", DASH + "CareerDashboard.tsx → conceptPathTab"),
    ("Web + mobile", "Path tab", "'Kỹ năng cần xây' for the selected step", " | ".join(stage_groups(ex["stage_skills"])), "Sheet", "stage_skills", "4 groups split on '., ' → Entry=step 1, Mid=step 2, Senior=step 3, Executive=step 5; step 4 has none in the sheet", "careers[].legacyStages[].skills", "scripts/import-alljobs-sheet.py → STAGE_TO_RUNG"),
    ("Web", "Right rail", "Độ phù hợp ring (88%)", "user × career", "Computed", "riasec_R..C + user scores", "Pearson correlation of the 6 RIASEC scores → (r+1)/2 × 100. Label: ≥75 cao, ≥60 khá, ≥40 khám phá", "—", DASH + "profile-match.ts"),
    ("Web", "Right rail", "4 bars: Kỹ năng / Kinh nghiệm / Sở thích / Giá trị cá nhân", "88 / 88 / 88 / 83", "Computed (display only)", "—", "All four are views of the same fit score; 'Giá trị cá nhân' = score × 0.94. Not separate measurements — review before production", "—", DASH + "CareerDashboard.tsx → fitBreakdown"),
    ("Web", "Right rail", "Nghề đề xuất list ('Đề xuất')", "9 careers", "Backend + crosswalk", "name_en / name_vi", "latest_survey.job_top_5, then job_zone_rec zone_5, zone_4, zone_3 → matched to sheet names, else via sheet 'Report_crosswalk'", "careers[]", DASH + "report-recommendations.ts"),
    ("Web", "Right rail", "'Tất cả' list", f"{len(careers)} careers", "Sheet", "all rows", "Every career, narrowed by search and filters", "careers[]", DASH + "CareerDashboard.tsx → conceptRecs"),
    ("Web", "Right rail", "Note 'Nghiên cứu và … là những điểm giao…'", "top shared traits", "Computed", "riasec_R..C + user scores", "3 RIASEC letters with the highest user × career product", "—", DASH + "profile-match.ts → topSharedTraits"),
    ("Mobile", "Mobile card", "Category chip on the card art", ex["category_job_group"], "Sheet", "category_job_group", "", "careers[].jobGroup", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "'#127 · RIASEC IS'", f"#{ex['id']:03d} · RIASEC {ex['riasec_code']}", "Sheet", "id, riasec_code", "", "careers[].id / riasec.code", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "Title + summary", ex["name_vi"], "Sheet", "name_vi / name_en, day_to_day", "", "careers[].name / dayToDay", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "Lương tại Việt Nam", vnd(ex["salary_vietnam"]), "Sheet", "salary_vietnam", "Same split as the web strip", "careers[].salary.vietnam", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "Hạng phổ biến '#52 · Đang tăng'", f"#{ex['popularity_rank']} · {TREND_VI[ex['trend_status']]}", "Sheet", "popularity_rank, trend_status", "", "careers[].trend", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "'Nghề liên quan' deck mode", "8 careers", "Sheet", "similar_careers", "Deck shows the current career + its similar careers", "careers[].similarCareerSlugs", DASH + "MobileCareerDeck.tsx"),
    ("Mobile", "Mobile card", "Detail sheet salary (US)", ex["salary_united_states"], "Sheet", "salary_united_states", "Shown only in the mobile detail sheet (Tổng quan)", "careers[].salary.unitedStates", DASH + "CareerDashboard.tsx → OverviewPanel"),
    ("Web + mobile", "Backend (per user)", "Lưu nghề / heart", "bookmark", "Backend", "slug, name_vi", "GET/POST /api/auth/profile/saved-careers/ {career_slug, career_name}", "—", DASH + "saved-careers-repository.ts"),
]

BACKEND = [
    ("GET", "/api/auth/profile/", "latest_survey.fish_type", "Fish theme (colours, mascot) and fish chip", "Salmon / Shark / Dolphin / Clownfish / Swordfish / Seahorse"),
    ("GET", "/api/auth/profile/", "latest_survey.riasec_scores", "Radar 'Hồ sơ của bạn', Độ phù hợp, shared-trait note", "R I A S E C, 1–5 scale"),
    ("GET", "/api/auth/profile/", "latest_survey.job_top_5", "Nghề đề xuất (first, in order)", "Names from backend survey/mappings/prod_job_filtered.csv"),
    ("GET", "/api/auth/profile/", "latest_survey.job_zone_rec.zone_5 / zone_4 / zone_3", "Nghề đề xuất (after job_top_5)", "Duplicates removed"),
    ("GET", "/api/auth/profile/", "latest_survey.id", "'Báo cáo chi tiết' button", "Survey result UUID"),
    ("POST", "/api/survey/report/generate/<id>/", "body {language}; response {pdf_url}", "'Báo cáo chi tiết' opens pdf_url", "Existing endpoint; ~2 min the first time (AI), instant afterwards"),
    ("GET", "/api/auth/profile/saved-careers/", "{slugs[], items[]}", "Heart state, saved marks in lists", "Added by this work: user.SavedCareer model + migration (backend/ folder)"),
    ("POST", "/api/auth/profile/saved-careers/", "body {career_slug, career_name}; response {saved, career_slug}", "Lưu nghề toggle", "Per-user, unique (user, career_slug)"),
]

wb = Workbook()


def style_header(cell, fill="2B1A3D"):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = BORDER


def body(cell, bold=False, color=INK):
    cell.font = Font(name=FONT, size=10, bold=bold, color=color)
    cell.alignment = WRAP
    cell.border = BORDER


# README
ws = wb.active
ws.title = "README"
lines = [
    ("QuickFish Career Path V2 — data map", True),
    ("Which data feeds which part of the V2 UI (web + mobile). Generated by scripts/build-data-map.py; re-run it after the sheet changes.", False),
    ("", False),
    ("Data flow", True),
    ("1. data/ALLJOBs_slim.xlsx (132 careers, 47 columns) → scripts/import-alljobs-sheet.py → public/data/career-dashboard.v1.json → the dashboard fetches it at runtime.", False),
    ("2. Per-user data (fish, RIASEC scores, recommended careers, detail report, bookmarks) comes from the QuickFish backend — see sheet 'Backend_API'.", False),
    ("3. Nothing outside these two sources is shown. Fields the sheet does not have stay empty in the UI.", False),
    ("", False),
    ("Sheets", True),
    ("UI_Map — one row per UI element: where it is, an example value (career: Nhà Khoa học Xã hội), the sheet column(s), the rule applied, the JSON field and the code that renders it.", False),
    ("Careers_by_UI — the ALLJOBs data re-ordered by UI section (row 1 = section, row 2 = label in the UI, row 3 = original sheet column).", False),
    ("Backend_API — per-user fields and endpoints.", False),
    ("Report_crosswalk — survey report job names that differ from sheet names, and the career they open. Hand-made; please review.", False),
    ("ALLJOBs_source — the original sheet, unchanged.", False),
    ("", False),
    ("Review flags", True),
    ("• Right-rail fit bars: the 4 bars are display views of one fit score ('Giá trị cá nhân' = score × 0.94). Not separate measurements.", False),
    ("• stage_skills has 4 groups for 5 path steps: step 4 has no skills in the sheet.", False),
    ("• day_to_day is English only; the Vietnamese UI shows it as is.", False),
    ("• Sheet column 'departments' is per career, not per major.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = ws.cell(row=i, column=1, value=text)
    cell.font = Font(name=FONT, size=14 if i == 1 else 10, bold=bold, color=PLUM if bold else INK)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

# UI_Map
ws = wb.create_sheet("UI_Map")
cols = ["#", "Screen", "UI area", "UI element", f"Example ({ex['name_vi']})", "Source", "Sheet column(s)", "Rule / transform", "JSON field (career-dashboard.v1.json)", "Code"]
widths = [5, 13, 20, 38, 42, 16, 30, 52, 30, 46]
for c, (name, width) in enumerate(zip(cols, widths), start=1):
    style_header(ws.cell(row=1, column=c, value=name))
    ws.column_dimensions[get_column_letter(c)].width = width
for r, row in enumerate(UI_MAP, start=2):
    values = (r - 1,) + row
    fill = PatternFill("solid", fgColor=SECTION_FILL.get(row[1], "FFFFFF"))
    for c, value in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=value)
        body(cell, bold=c in (4, 7))
        if c == 3:
            cell.fill = fill
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(UI_MAP) + 1}"

# Careers_by_UI
ws = wb.create_sheet("Careers_by_UI")
layout = [
    ("Identity & URL", [("ID (#)", "id"), ("Slug (URL key)", "slug")]),
    ("Hero (header)", [("Tên nghề (VI)", "name_vi"), ("Career name (EN)", "name_en"), ("Nhóm nghề", "category_job_group"), ("RIASEC pill", "riasec_code"), ("Mô tả (day to day)", "day_to_day")]),
    ("Metric strip", [("Lương khởi điểm (VN)", "salary_vietnam"), ("Hạng phổ biến", "popularity_rank"), ("Xu hướng", "trend_status"), ("Kỹ năng nổi bật", "skill_1"), ("% nhà tuyển dụng", "skill_1_employers_pct")]),
    ("Overview tab", [("Radar R", "riasec_R"), ("Radar I", "riasec_I"), ("Radar A", "riasec_A"), ("Radar S", "riasec_S"), ("Radar E", "riasec_E"), ("Radar C", "riasec_C")]),
    ("Skills tab", [x for i in range(2, 9) for x in ((f"Kỹ năng {i}", f"skill_{i}"), (f"% {i}", f"skill_{i}_employers_pct"))]),
    ("Majors tab", [("Ngành 1", "major_1"), ("Ngành 2", "major_2"), ("Ngành 3", "major_3"), ("Khối ngành", "departments"), ("Môn cốt lõi", "core_subjects"), ("Khối xét tuyển", "exam_blocks")]),
    ("Path tab", [(f"Bậc {i}", f"path_{i}") for i in range(1, 6)] + [("Kỹ năng bậc 1 (Entry)", "stage_skills#0"), ("Kỹ năng bậc 2 (Mid)", "stage_skills#1"), ("Kỹ năng bậc 3 (Senior)", "stage_skills#2"), ("Kỹ năng bậc 5 (Executive)", "stage_skills#3")]),
    ("Related careers", [("Nghề tương tự (slugs)", "similar_careers")]),
    ("Mobile card", [("Lương Mỹ (detail sheet)", "salary_united_states")]),
]
col = 1
for section, fields in layout:
    start = col
    fill = PatternFill("solid", fgColor=SECTION_FILL[section])
    for label, source in fields:
        style_header(ws.cell(row=2, column=col, value=label), fill="4A3560")
        src = ws.cell(row=3, column=col, value=source.replace("#", " group "))
        src.font = Font(name=FONT, size=9, italic=True, color="667085")
        src.fill = fill
        src.border = BORDER
        for r, career in enumerate(careers, start=4):
            if source.startswith("stage_skills#"):
                groups = stage_groups(career["stage_skills"])
                idx = int(source.split("#")[1])
                value = groups[idx] if idx < len(groups) else None
            else:
                value = career[source]
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=source in ("day_to_day", "similar_careers"))
        ws.column_dimensions[get_column_letter(col)].width = 44 if source in ("day_to_day", "similar_careers") else 18 if not source.startswith(("riasec_", "skill_1_", "popularity", "id")) else 10
        col += 1
    ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
    head = ws.cell(row=1, column=start, value=section)
    head.font = Font(name=FONT, bold=True, size=11, color=INK)
    head.fill = fill
    head.alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "D4"
ws.auto_filter.ref = f"A3:{get_column_letter(col - 1)}{len(careers) + 3}"
ws.row_dimensions[2].height = 32

# Backend_API
ws = wb.create_sheet("Backend_API")
cols = ["Method", "Endpoint", "Field / payload", "Used for (UI)", "Notes"]
for c, (name, width) in enumerate(zip(cols, [9, 40, 46, 46, 56]), start=1):
    style_header(ws.cell(row=1, column=c, value=name))
    ws.column_dimensions[get_column_letter(c)].width = width
for r, row in enumerate(BACKEND, start=2):
    for c, value in enumerate(row, start=1):
        body(ws.cell(row=r, column=c, value=value), bold=c == 3)
ws.freeze_panes = "A2"

# Report_crosswalk
ws = wb.create_sheet("Report_crosswalk")
pairs = re.findall(r'^\s*"([^"]+)": "([a-z0-9-]+)",', CROSSWALK.read_text(encoding="utf-8"), re.M)
cols = ["Report job name (normalised)", "Opens career (slug)", "Career name (VI)", "Career name (EN)"]
for c, (name, width) in enumerate(zip(cols, [36, 52, 40, 52]), start=1):
    style_header(ws.cell(row=1, column=c, value=name))
    ws.column_dimensions[get_column_letter(c)].width = width
for r, (report_name, slug) in enumerate(pairs, start=2):
    career = by_slug.get(slug, {})
    for c, value in enumerate((report_name, slug, career.get("name_vi"), career.get("name_en")), start=1):
        body(ws.cell(row=r, column=c, value=value))
note_row = len(pairs) + 3
ws.cell(row=note_row, column=1, value="Report names that match a sheet name exactly need no row here. Names not listed and not matching (e.g. Teacher, Therapist) show as 'not in the dashboard'.").font = Font(name=FONT, italic=True, size=9, color="667085")
ws.freeze_panes = "A2"

# ALLJOBs_source
ws = wb.create_sheet("ALLJOBs_source")
for c, name in enumerate(headers, start=1):
    style_header(ws.cell(row=1, column=c, value=name))
    ws.column_dimensions[get_column_letter(c)].width = 16
for r, row in enumerate(rows[1:], start=2):
    for c, value in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=value).font = Font(name=FONT, size=10)
ws.freeze_panes = "C2"

wb.save(OUT)
print(f"Wrote {OUT.relative_to(ROOT)}: {len(UI_MAP)} UI elements, {len(careers)} careers, {len(pairs)} crosswalk rows")
